import os
from PyPDF2 import PdfReader
from docx import Document
from pptx import Presentation
from config import Config


def allowed_file(filename):
    """Check if file extension is allowed."""
    return '.' in filename and \
           filename.rsplit('.', 1)[1].lower() in Config.ALLOWED_EXTENSIONS


def extract_text_from_pdf(file_path):
    """Extract text from PDF file."""
    try:
        reader = PdfReader(file_path)
        text = ""
        for page in reader.pages:
            page_text = page.extract_text()
            if page_text:
                text += page_text + "\n"
        return text.strip()
    except Exception as e:
        return f"Error extracting PDF text: {str(e)}"


def extract_text_from_docx(file_path):
    """Extract text from DOCX file."""
    try:
        doc = Document(file_path)
        text = ""
        for paragraph in doc.paragraphs:
            text += paragraph.text + "\n"
        # Also extract text from tables
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    text += cell.text + "\t"
                text += "\n"
        return text.strip()
    except Exception as e:
        return f"Error extracting DOCX text: {str(e)}"


def extract_text_from_pptx(file_path):
    """Extract text from PPTX file."""
    try:
        prs = Presentation(file_path)
        text = ""
        for slide_num, slide in enumerate(prs.slides, 1):
            text += f"\n--- Slide {slide_num} ---\n"
            for shape in slide.shapes:
                if hasattr(shape, "text"):
                    text += shape.text + "\n"
        return text.strip()
    except Exception as e:
        return f"Error extracting PPTX text: {str(e)}"


def extract_text_from_txt(file_path):
    """Extract text from TXT file."""
    try:
        with open(file_path, 'r', encoding='utf-8') as f:
            return f.read().strip()
    except UnicodeDecodeError:
        try:
            with open(file_path, 'r', encoding='latin-1') as f:
                return f.read().strip()
        except Exception as e:
            return f"Error reading TXT file: {str(e)}"
    except Exception as e:
        return f"Error reading TXT file: {str(e)}"


def extract_text(file_path):
    """Extract text from file based on extension."""
    ext = file_path.rsplit('.', 1)[1].lower()

    if ext == 'pdf':
        return extract_text_from_pdf(file_path)
    elif ext == 'docx':
        return extract_text_from_docx(file_path)
    elif ext == 'pptx':
        return extract_text_from_pptx(file_path)
    elif ext == 'txt':
        return extract_text_from_txt(file_path)
    else:
        return "Unsupported file type"


def get_file_size_mb(file_path):
    """Get file size in megabytes."""
    size_bytes = os.path.getsize(file_path)
    return round(size_bytes / (1024 * 1024), 2)


def ensure_upload_directory(brand_name):
    """Ensure upload directory exists for a brand."""
    brand_dir = os.path.join(Config.UPLOAD_FOLDER, brand_name.replace(' ', '_'))
    os.makedirs(brand_dir, exist_ok=True)
    return brand_dir


def save_uploaded_file(file, brand_name):
    """Save uploaded file and return the file path."""
    brand_dir = ensure_upload_directory(brand_name)
    filename = secure_filename(file.filename)

    # Handle duplicate filenames
    base_name, ext = os.path.splitext(filename)
    counter = 1
    while os.path.exists(os.path.join(brand_dir, filename)):
        filename = f"{base_name}_{counter}{ext}"
        counter += 1

    file_path = os.path.join(brand_dir, filename)
    file.save(file_path)
    return file_path, filename


def secure_filename(filename):
    """Make filename safe for storage."""
    # Remove path separators and keep only safe characters
    keepchars = (' ', '.', '_', '-')
    filename = "".join(c for c in filename if c.isalnum() or c in keepchars).rstrip()
    # Replace spaces with underscores
    filename = filename.replace(' ', '_')
    return filename if filename else 'unnamed_file'


def delete_file(file_path):
    """Delete a file if it exists."""
    try:
        if os.path.exists(file_path):
            os.remove(file_path)
            return True
    except Exception:
        pass
    return False
